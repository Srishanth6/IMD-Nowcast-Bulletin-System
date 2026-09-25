"""
IMD Nowcast Bulletin — single-page IMD format.

Combines Person 1 radar and Person 2 warning-map outputs into a Word
bulletin that follows the required one-page Telangana nowcast layout.
Times come from the live IMD nowcast page, not hard-coded values.
"""

from __future__ import annotations

import html as html_lib
import io
import os
import re
import shutil
import sys
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urljoin
from zoneinfo import ZoneInfo

import requests
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

from download_warning_map import URL as IMD_NOWCAST_URL, extract_svg
from integrate_bulletin import OUTPUT_PATH as COMPOSITE_PATH
from integrate_bulletin import main as build_composite_image


ROOT = Path(__file__).resolve().parent

RADAR_PATH = ROOT / "radar_download" / "radar_images" / "latest_radar.png"
WARNING_MAP_PATH = ROOT / "latest_warning_map.png"
WARNING_SVG_PATH = ROOT / "latest_warning_map.svg"
HEADER_PATH = ROOT / "assets" / "imd_header.jpg"
PAGE2_IMPACTS_PATH = ROOT / "assets" / "expected_impacts_moderate.jpg"
PAGE3_IMPACTS_PATH = ROOT / "assets" / "expected_impacts_heavy_rain.jpg"
PAGE4_IMPACTS_PATH = ROOT / "assets" / "expected_impacts_very_severe.jpg"
OUTPUT_PATH = ROOT / "IMD_Nowcast_Bulletin.docx"
SERVED_BULLETIN_NAME = "telangana_nowcast.docx"
DOWNLOAD_PHP_PATH = ROOT / "dist_nowcast3.php"
DEFAULT_WX_DIR = Path("/var/www/html/tlng/wx")
DOWNLOAD_BULLETIN_URL = urljoin(IMD_NOWCAST_URL, "dist_nowcast3.php")

try:
    IST = ZoneInfo("Asia/Kolkata")
except Exception:
    IST = timezone(timedelta(hours=5, minutes=30))
NOWCAST_VALIDITY = timedelta(hours=3)

ORANGE_FILL = "FFA500"
YELLOW_FILL = "FFFF00"
NAVY = RGBColor(0x0C, 0x2F, 0x52)
BLUE = RGBColor(0x20, 0x5B, 0x91)
INK = RGBColor(0x24, 0x34, 0x45)
GOLD = RGBColor(0xE2, 0xA2, 0x2C)
PALE_FILL = "EBF3FA"
BORDER_BLUE = "205B91"
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

LEVELS = (
    ("red", "Red Warning(Take Action)", "FF0000"),
    ("orange", "Orange Warning(Be Prepared)", ORANGE_FILL),
    ("yellow", "Yellow Warning(Be Updated)", YELLOW_FILL),
)

ALWAYS_SHOW_LEVELS = {"orange", "yellow"}


class BulletinInputError(FileNotFoundError):
    """Raised when a required bulletin input image is missing."""


def datetimes_from_imd(page_fields: dict):
    """Use the issue time published on the latest IMD page, not the local clock."""
    date_db = page_fields.get("date_db")
    toi = page_fields.get("time_toi_db")
    valid_text = page_fields.get("valid_upto_db")
    if not date_db or not toi:
        raise BulletinInputError(
            "The latest IMD page did not include Date (DB) and Time TOI (DB)."
        )
    issued = datetime.strptime(f"{date_db} {toi}", "%Y-%m-%d %H:%M:%S").replace(tzinfo=IST)
    if valid_text:
        valid = datetime.strptime(f"{date_db} {valid_text}", "%Y-%m-%d %H:%M:%S").replace(tzinfo=IST)
        if valid <= issued:
            valid += timedelta(days=1)
    else:
        valid = issued + NOWCAST_VALIDITY
    return issued, valid


def _has_telugu(text: str) -> bool:
    return bool(re.search(r"[\u0C00-\u0C7F]", text))


def parse_nowcast_page(page_html: str) -> dict:
    """Extract fields that exist on dist_nowcast.php (map page)."""
    prefix = page_html.split("<!DOCTYPE", 1)[0]
    fields = {
        "date_db": None,
        "time_toi_db": None,
        "valid_upto_db": None,
        "official_bulletin_href": None,
        "map_fill_counts": Counter(),
    }
    date_match = re.search(r"Date \(DB\):\s*([0-9]{4}-[0-9]{2}-[0-9]{2})", prefix)
    toi_match = re.search(r"Time TOI \(DB\):\s*([0-9]{1,2}:[0-9]{2}:[0-9]{2})", prefix)
    valid_match = re.search(r"Valid Up To \(DB\):\s*([0-9]{1,2}:[0-9]{2}:[0-9]{2})", prefix)
    if date_match:
        fields["date_db"] = date_match.group(1)
    if toi_match:
        fields["time_toi_db"] = toi_match.group(1)
    if valid_match:
        fields["valid_upto_db"] = valid_match.group(1)

    href_match = re.search(
        r"<a\s+href=['\"]([^'\"]+)['\"][^>]*>\s*Download Nowcast Bulletin",
        page_html,
        flags=re.I,
    )
    if href_match:
        fields["official_bulletin_href"] = href_match.group(1)

    svg = extract_svg(page_html)
    fills = re.findall(
        r"<path[^>]*fill=['\"](green|yellow|orange|red)['\"]",
        svg,
        flags=re.I,
    )
    fields["map_fill_counts"] = Counter(color.lower() for color in fills)
    return fields


def parse_official_bulletin_docx(docx_bytes: bytes) -> dict:
    """Extract Orange/Yellow/Red English and Telugu text from IMD's Word bulletin."""
    sections = {}
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as archive:
        xml = archive.read("word/document.xml")
    root = ET.fromstring(xml)
    runs = []
    for node in root.iter(f"{W_NS}t"):
        if node.text:
            runs.append(html_lib.unescape(node.text).strip())

    heading_keys = {
        "orange warning": "orange",
        "yellow warning": "yellow",
        "red warning": "red",
    }
    index = 0
    while index < len(runs):
        compact = re.sub(r"\s+", "", runs[index]).lower()
        matched_key = None
        for prefix, key in heading_keys.items():
            if compact.startswith(prefix.replace(" ", "")):
                matched_key = key
                break
        if matched_key is None:
            index += 1
            continue
        english = ""
        telugu = ""
        if index + 1 < len(runs):
            english = runs[index + 1]
        if index + 2 < len(runs) and _has_telugu(runs[index + 2]):
            telugu = runs[index + 2]
            index += 3
        else:
            index += 2
        sections[matched_key] = {"english": english, "telugu": telugu}
    return sections


def fetch_nowcast_page() -> str:
    page = requests.get(
        IMD_NOWCAST_URL,
        params={"_": str(int(datetime.now(IST).timestamp()))},
        headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
        timeout=30,
    )
    page.raise_for_status()
    html_path = ROOT / "imd_response.html"
    html_path.write_text(page.text, encoding="utf-8")
    return page.text


def load_imd_warning_content(page_html: str) -> dict:
    page_fields = parse_nowcast_page(page_html)

    print("Fields available on dist_nowcast.php:")
    print(f"  Date (DB): {page_fields['date_db']}")
    print(f"  Time TOI (DB): {page_fields['time_toi_db']}")
    print(f"  Valid Up To (DB): {page_fields['valid_upto_db']}")
    print(f"  Map polygon fills: {dict(page_fields['map_fill_counts'])}")
    print(f"  Official bulletin link: {page_fields['official_bulletin_href']}")
    print("  English/Telugu warning sentences: not present on this page")
    print("  District name attributes: not present in the SVG")

    warning_sections = {}
    href = page_fields["official_bulletin_href"]
    if href:
        bulletin_url = urljoin(IMD_NOWCAST_URL, href)
        print("Fetching official IMD warning text from:")
        print(f"  {bulletin_url}")
        bulletin = requests.get(bulletin_url, timeout=60)
        bulletin.raise_for_status()
        warning_sections = parse_official_bulletin_docx(bulletin.content)
        print("Fields available in the official IMD Word bulletin:")
        for key in ("orange", "yellow", "red"):
            section = warning_sections.get(key)
            if section:
                print(f"  {key} English: {section['english'][:90]}...")
                print(f"  {key} Telugu: present" if section["telugu"] else f"  {key} Telugu: absent")
            else:
                print(f"  {key}: not present")

    return {
        "page_fields": page_fields,
        "warning_sections": warning_sections,
    }


def format_issue_line(issued: datetime) -> str:
    return f"TIME OF ISSUE: {issued.strftime('%Y-%m-%d')} ({issued.strftime('%H:%M:%S')} Hrs IST)"


def format_valid_line(valid: datetime) -> str:
    return f"Valid upto: ({valid.strftime('%H:%M:%S')} Hrs IST)"


def format_bottom_timestamp(issued: datetime, valid: datetime) -> str:
    return (
        f"Date: {issued.strftime('%Y-%m-%d')}  |  "
        f"Time of Issue: {issued.strftime('%H:%M:%S')}  |  "
        f"Valid Up To: {valid.strftime('%H:%M:%S')}"
    )


def set_run_font(run, name="Arial", size=11, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    r_fonts = rpr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        rpr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), name)
    if color is not None:
        run.font.color.rgb = color


def set_cell_width(cell, width):
    cell.width = width


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_borders(cell, color=BORDER_BLUE, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    tc_pr.append(borders)


def set_cell_margins(cell, top=100, bottom=100, start=140, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def add_paragraph(document, text="", *, align="left", space_after=4, space_before=0):
    paragraph = document.add_paragraph()
    alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.0
    if text:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def picture_size(path: Path, max_width_in: float, max_height_in: float):
    with Image.open(path) as image:
        width_px, height_px = image.size
    aspect = width_px / height_px
    width = max_width_in
    height = width / aspect
    if height > max_height_in:
        height = max_height_in
        width = height * aspect
    return Inches(width), Inches(height)


def require_input(path: Path, label: str):
    if not path.is_file():
        raise BulletinInputError(
            f"{label} not found:\n  {path.relative_to(ROOT).as_posix()}\n"
            "Generate this file with the existing module before creating the bulletin."
        )
    print(f"Loaded {label}: {path.relative_to(ROOT).as_posix()}")
    return path


def add_page_header(section):
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("INDIA METEOROLOGICAL DEPARTMENT")
    set_run_font(run, size=10, bold=True, color=NAVY)


def add_datetime_table(document, issued: datetime, valid: datetime):
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell_width = Inches(2.366)
    values = (
        ("BULLETIN DATE", issued.strftime("%Y-%m-%d")),
        ("TIME OF ISSUE", issued.strftime("%H:%M:%S")),
        ("VALID UP TO", valid.strftime("%H:%M:%S")),
    )
    for cell, (label, value) in zip(table.rows[0].cells, values):
        set_cell_width(cell, cell_width)
        shade_cell(cell, PALE_FILL)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(2)
        label_run = paragraph.add_run(label)
        set_run_font(label_run, size=8.5, bold=True, color=BLUE)
        label_run.add_break()
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=11, bold=True, color=INK)
    return table


def build_document(radar_path: Path, map_path: Path, warning_data: dict, issued: datetime, valid: datetime):
    print(format_issue_line(issued))
    print(format_valid_line(valid))
    print(format_bottom_timestamp(issued, valid))
    require_input(radar_path, "Hyderabad radar image")
    require_input(map_path, "Telangana warning map")

    print("Building IMD-style bulletin visual...")
    build_composite_image(issued=issued, valid=valid)
    composite_path = require_input(COMPOSITE_PATH, "integrated bulletin visual")

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.25)
    add_page_header(section)

    title = add_paragraph(document, align="center", space_before=6, space_after=2)
    title_run = title.add_run("NOWCAST BULLETIN")
    set_run_font(title_run, size=24, bold=True, color=NAVY)

    state = add_paragraph(document, align="center", space_before=0, space_after=12)
    state_run = state.add_run("TELANGANA")
    set_run_font(state_run, size=15, bold=True, color=BLUE)

    add_datetime_table(document, issued, valid)
    add_paragraph(document, space_after=0)

    visual = add_paragraph(document, align="center", space_before=4, space_after=10)
    width, height = picture_size(composite_path, 7.1, 4.9)
    visual.add_run().add_picture(str(composite_path), width=width, height=height)

    separator = add_paragraph(document, align="center", space_before=0, space_after=5)
    sep_run = separator.add_run("_" * 64)
    set_run_font(sep_run, size=7, color=GOLD)

    source = add_paragraph(document, align="center", space_after=0)
    source_run = source.add_run(
        "Source: India Meteorological Department | Telangana nowcast products"
    )
    set_run_font(source_run, size=8.5, color=INK)

    fill_counts = warning_data.get("page_fields", {}).get("map_fill_counts") or Counter()
    print(f"Map polygon fills used in the visual: {dict(fill_counts)}")
    return document


def nowcast_wx_directories() -> list[Path]:
    """Directories that Apache/PHP can use for Download Nowcast Bulletin."""
    directories: list[Path] = []
    seen: set[Path] = set()
    env_dir = os.environ.get("IMD_WX_DIR", "").strip()
    candidates = []
    if env_dir:
        candidates.append(Path(env_dir))
    candidates.append(DEFAULT_WX_DIR)
    candidates.append(ROOT)
    for candidate in candidates:
        resolved = candidate.resolve()
        if resolved in seen:
            continue
        seen.add(resolved)
        directories.append(candidate)
    return directories


def _copy_if_needed(source: Path, destination: Path) -> Path | None:
    if not source.is_file():
        return None
    if destination.exists() and destination.resolve() == source.resolve():
        return destination
    try:
        shutil.copy2(source, destination)
    except OSError:
        return None
    return destination


def publish_generated_bulletin(docx_path: Path) -> list[Path]:
    """Copy the generated bulletin (and download PHP) into the IMD wx folder."""
    published: list[Path] = []
    seen: set[Path] = set()
    for directory in nowcast_wx_directories():
        if not directory.is_dir() or not os.access(directory, os.W_OK):
            continue
        for name in (OUTPUT_PATH.name, SERVED_BULLETIN_NAME):
            copied = _copy_if_needed(docx_path, directory / name)
            if copied is not None:
                key = copied.resolve()
                if key not in seen:
                    seen.add(key)
                    published.append(copied)
        php_copy = _copy_if_needed(DOWNLOAD_PHP_PATH, directory / DOWNLOAD_PHP_PATH.name)
        if php_copy is not None:
            key = php_copy.resolve()
            if key not in seen:
                seen.add(key)
                published.append(php_copy)
    return published


def live_download_matches(docx_path: Path) -> bool:
    """True when the IP Download Nowcast Bulletin link returns this generated file."""
    try:
        response = requests.get(
            DOWNLOAD_BULLETIN_URL,
            timeout=60,
            headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
        )
        response.raise_for_status()
    except Exception as error:
        print(f"WARNING: could not check Download Nowcast Bulletin: {error}")
        return False
    return response.content == docx_path.read_bytes()


def main():
    print("=" * 60)
    print(" IMD NOWCAST BULLETIN")
    print("=" * 60)
    os.chdir(ROOT)
    print("Fetching latest IMD data...")

    try:
        page_html = fetch_nowcast_page()
    except Exception as error:
        print(f"ERROR: could not fetch the latest IMD nowcast page: {error}")
        return 1

    try:
        from radar_download.radar_scraper import download_latest_radar

        print("Downloading latest radar image...")
        radar_file = download_latest_radar()
    except Exception as error:
        print(f"ERROR: could not download the latest radar image: {error}")
        return 1
    if not radar_file or not Path(radar_file).is_file():
        print("ERROR: the latest radar image was not downloaded.")
        return 1
    radar_path = Path(radar_file)
    print(f"Latest radar saved: {radar_path.as_posix()}")

    try:
        from download_warning_map import save_warning_map

        print("Downloading latest warning map...")
        map_path = save_warning_map(page_html)
    except Exception as error:
        print(f"ERROR: could not download the latest warning map: {error}")
        return 1
    print(f"Latest warning map saved: {Path(map_path).as_posix()}")

    try:
        warning_data = load_imd_warning_content(page_html)
        issued, valid = datetimes_from_imd(warning_data["page_fields"])
    except Exception as error:
        print(f"ERROR: could not extract the IMD issue time: {error}")
        return 1

    print(
        "Issue time extracted: "
        f"{issued.strftime('%Y-%m-%d')} ({issued.strftime('%H:%M:%S')} Hrs IST)"
    )
    print("Generating final bulletin...")

    try:
        document = build_document(radar_path, map_path, warning_data, issued, valid)
        document.save(OUTPUT_PATH)
        published = publish_generated_bulletin(OUTPUT_PATH)
    except Exception as error:
        print(f"ERROR: could not generate the bulletin: {error}")
        return 1

    print("Bulletin generated successfully!")
    print(f"Saved: {OUTPUT_PATH.name}")
    if published:
        print("Published for Download Nowcast Bulletin:")
        for path in published:
            print(f"  {path.as_posix()}")
        print(f"Download URL: {DOWNLOAD_BULLETIN_URL}")
        if live_download_matches(OUTPUT_PATH):
            print("Download Nowcast Bulletin is serving this generated bulletin.")
        else:
            print(
                "WARNING: the IP download link is still serving the previous "
                "server-generated bulletin. Run this script on the IMD Apache "
                "host (or set IMD_WX_DIR to /var/www/html/tlng/wx) so "
                "dist_nowcast3.php and IMD_Nowcast_Bulletin.docx land in that folder."
            )
    else:
        print(
            "WARNING: could not copy the bulletin into the IMD wx folder. "
            "Set IMD_WX_DIR to /var/www/html/tlng/wx on the Apache server."
        )
    return 0


if __name__ == "__main__":
    sys.exit(main())
